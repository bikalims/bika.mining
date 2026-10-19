# -*- coding: utf-8 -*-
import unittest
import zipfile
from io import BytesIO
from copy import copy
from openpyxl import load_workbook
from bika.mining.resultsxlsx import TEMPLATE, render_results
from senaite.core.tests.layers import BASE_LAYER_FIXTURE
from plone.app.testing import PloneSandboxLayer, IntegrationTesting, applyProfile
import bika.mining
from bika.mining.interfaces import IBikaMiningLayer
from zope.interface import alsoProvides
from zope.component import getMultiAdapter
from bika.lims.interfaces import IClient
from zope.interface import implementer


class TestResultsXLSX(unittest.TestCase):
    def export(self, rows):
        return load_workbook(BytesIO(render_results(rows)))

    def test_fixed_columns_blank_cells_and_zero(self):
        workbook = self.export([
            {"metadata": {"A": "Sample-1", "M": 0},
             "results": {"Moisture (%)": "0", "NH4 (ug/gU)": "<0.1"}},
            {"metadata": {"A": "Sample-2"},
             "results": {"U3O8 solids (ppm)": "12.30"}},
        ])
        sheet = workbook["Reporting"]
        self.assertEqual(sheet["Y4"].value, "0")
        self.assertEqual(sheet["DQ4"].value, "<0.1")
        self.assertEqual(sheet["P5"].value, "12.30")
        self.assertIsNone(sheet["P4"].value)
        self.assertIsNone(sheet["Y5"].value)
        self.assertIsNone(sheet["B4"].value)
        self.assertEqual(sheet["M4"].value, 0)
        self.assertEqual(sheet.max_row, 5)
        self.assertEqual(sheet.max_column, 121)
        self.assertEqual(sheet.auto_filter.ref, "B3:DQ5")

    def test_template_formatting_and_headers_preserved(self):
        original = load_workbook(TEMPLATE)["Reporting"]
        sheet = self.export([{"metadata": {}, "results": {}}])["Reporting"]
        for row in original.iter_rows(min_row=1, max_row=3):
            for cell in row:
                expected = None if cell.coordinate in ('A1', 'C1', 'E1', 'G1') else cell.value
                self.assertEqual(sheet[cell.coordinate].value, expected)
                self.assertEqual(sheet[cell.coordinate]._style, cell._style)
        self.assertEqual(list(sheet.merged_cells.ranges),
                         list(original.merged_cells.ranges))
        for row in sheet.iter_rows(min_row=4, max_row=4):
            for cell in row:
                self.assertEqual(cell.font.name, 'Arial')
                self.assertEqual(cell.font.color.rgb, 'FF000000')
                self.assertFalse(cell.font.italic)
                self.assertFalse(cell.font.bold)
                source_cell = original[cell.coordinate]
                self.assertEqual(copy(cell.alignment), copy(source_cell.alignment))
                self.assertEqual(copy(cell.fill), copy(source_cell.fill))
                self.assertEqual(copy(cell.border), copy(source_cell.border))
                self.assertEqual(cell.number_format, source_cell.number_format)
        self.assertEqual(sheet.column_dimensions["P"].width,
                         original.column_dimensions["P"].width)

    def test_chlorine_ppm_fills_both_fixed_sections_and_leaves_missing_blank(self):
        sheet = self.export([
            {"metadata": {}, "results": {"Cl (ppm)": "12"}},
            {"metadata": {}, "results": {}},
        ])["Reporting"]
        self.assertEqual(sheet["CB3"].value, "Cl (ppm)")
        self.assertEqual(sheet["DD3"].value, "Cl (ppm)")
        self.assertEqual(sheet["CB4"].value, "12")
        self.assertEqual(sheet["DD4"].value, "12")
        self.assertIsNone(sheet["CB5"].value)
        self.assertIsNone(sheet["DD5"].value)

    def test_more_samples_than_example_rows(self):
        rows = [{"metadata": {"A": "Sample-{}".format(i)},
                 "results": {"NH4 (ug/gU)": str(i)}} for i in range(12)]
        sheet = self.export(rows)["Reporting"]
        self.assertEqual(sheet["A15"].value, "Sample-11")
        self.assertEqual(sheet["DQ15"].value, "11")

    def test_v3_services_and_layout(self):
        destinations = {
            'AE': 'Granulometry (% passing 6 mm)', 'AF': 'Slurry SG (t/m3)',
            'AH': 'Solids (% w/w)', 'AI': 'Resin concentration (% v/v)',
            'AJ': 'Resin carry-over (g/L)', 'AK': 'Wet mass of cut (kg)',
            'AL': 'Dry mass of cut (kg)', 'AM': 'Mass recovered (kg)',
            'AN': 'Mass produced (t)', 'AO': 'Volume (m3)',
            'AT': 'pH', 'AU': 'ORP (mV)', 'AX': 'Dissolved oxygen (ppb)',
            'CN': 'Cycles of concentration (ratio)',
            'CO': 'Corrosion inhibitor residual (ppm)',
            'CQ': 'Resin working capacity (g U/L)',
            'CR': 'Bead size distribution (%)', 'CS': 'Whole-bead count (%)',
            'DF': 'CO3 (ug/gU)',
        }
        results = dict((name, str(index)) for index, name in
                       enumerate(sorted(destinations.values())))
        sheet = self.export([{'metadata': {}, 'results': results}])['Reporting']
        for col, name in destinations.items():
            self.assertEqual(sheet[col + '3'].value, name)
            self.assertEqual(sheet[col + '4'].value, results[name])
        self.assertEqual(sheet['A2'].value, 'LIMS Sample ID')
        self.assertEqual(sheet['AT2'].value, 'Field / water')
        for row in sheet.iter_rows(min_row=3, max_row=3, min_col=16, max_col=121):
            for cell in row:
                self.assertEqual(cell.alignment.textRotation, 90)

    def test_report_header_and_empty_export(self):
        header = {'A': 'Plant', 'C': 'COA26-0001 Plant.xlsx',
                  'E': '2026-10-18 09:15', 'G': 'Publisher'}
        sheet = load_workbook(BytesIO(render_results(
            [], report_header=header)))['Reporting']
        for col, value in header.items():
            self.assertEqual(sheet[col + '1'].value, value)
        self.assertEqual(sheet.max_row, 3)
        self.assertEqual(sheet.auto_filter.ref, 'B3:DQ3')

    def test_unknown_service_title_is_rejected(self):
        with self.assertRaises(ValueError):
            render_results([{"metadata": {}, "results": {"Unknown AS": "1"}}])

    def test_results_are_literal_text(self):
        sheet = self.export([{"metadata": {},
                              "results": {"Moisture (%)": "=1+1"}}])["Reporting"]
        self.assertEqual(sheet["Y4"].data_type, "s")

    def test_other_workbook_parts_are_unchanged(self):
        data = render_results([{"metadata": {}, "results": {}}])
        with zipfile.ZipFile(TEMPLATE) as source:
            with zipfile.ZipFile(BytesIO(data)) as output:
                for name in source.namelist():
                    if name not in ("xl/worksheets/sheet1.xml", "xl/workbook.xml",
                                    "xl/styles.xml"):
                        self.assertEqual(source.read(name), output.read(name))


class ResultsPublicationLayer(PloneSandboxLayer):
    defaultBases = (BASE_LAYER_FIXTURE,)

    def setUpZope(self, app, configurationContext):
        import plone.restapi
        import senaite.impress
        self.loadZCML(package=plone.restapi)
        self.loadZCML(package=senaite.impress)
        self.loadZCML(package=bika.mining)

    def setUpPloneSite(self, portal):
        # The SENAITE base fixture retains its earlier global registry.
        # Rebase the test site so this layer's add-on registrations are visible.
        from zope.component import getGlobalSiteManager
        portal.getSiteManager().__bases__ = (getGlobalSiteManager(),)
        from zope.component.hooks import setSite
        setSite(portal)
        applyProfile(portal, 'bika.mining:default')


RESULTS_TESTING = IntegrationTesting(bases=(ResultsPublicationLayer(),),
                                    name='MiningResults:Integration')


class TestResultsPublicationRegistration(unittest.TestCase):
    layer = RESULTS_TESTING

    def setUp(self):
        from zope.component.hooks import setSite
        from zope.component import getGlobalSiteManager
        self.layer['portal'].getSiteManager().__bases__ = (getGlobalSiteManager(),)
        setSite(self.layer['portal'])
        applyProfile(self.layer['portal'], 'bika.mining:default')

    def test_mining_publish_and_download_routes(self):
        from senaite.impress.publishview import PublishView
        from bika.mining.browser.resultsxlsx import ReportEmailView
        from bika.mining.browser.resultsxlsx import MiningAjaxPublishView

        @implementer(IClient)
        class Client(object):
            pass
        request = self.layer['request']
        from senaite.impress.interfaces import ISenaiteImpressLayer
        alsoProvides(request, ISenaiteImpressLayer)
        alsoProvides(request, IBikaMiningLayer)
        from zope.traversing.interfaces import BeforeTraverseEvent
        from zope.event import notify
        notify(BeforeTraverseEvent(self.layer['portal'], request))
        self.assertIsInstance(getMultiAdapter((Client(), request), name='publish'),
                              PublishView)
        self.assertIsInstance(getMultiAdapter((Client(), request), name='email'),
                              ReportEmailView)
        self.assertIsInstance(getMultiAdapter((Client(), request), name='ajax_publish'),
                              MiningAjaxPublishView)

    def test_sampling_methods_reuse_existing_records(self):
        from bika.lims import api
        from bika.mining.setuphandlers import setup_sampling_methods
        from plone.app.testing import setRoles, TEST_USER_ID
        setRoles(self.layer['portal'], TEST_USER_ID, ['Manager'])
        folder = api.get_senaite_setup().samplingdeviations
        original = dict((obj.Title(), api.get_uid(obj))
                        for obj in folder.objectValues())
        self.assertEqual(set(original), {'Grab', 'Composite', 'Profile/Inventory'})
        legacy = api.create(folder, 'SamplingDeviation', title=u'Existing method')
        setup_sampling_methods()
        self.assertEqual(folder.Title(), 'Sampling Methods')
        self.assertEqual(api.get_title(legacy), 'Existing method')
        for obj in folder.objectValues():
            if obj.Title() in original:
                self.assertEqual(api.get_uid(obj), original[obj.Title()])
        self.assertEqual(len(folder.objectValues()), 4)

    def test_publication_stores_xlsx_and_updates_sample_tat(self):
        from bika.lims import api
        from bika.mining.browser.resultsxlsx import MiningAjaxPublishView
        from bika.mining.browser.resultsxlsx import REPORT_TEMPLATE, refresh_xlsx
        from bika.mining.resultsxlsx import MIME_TYPE
        from plone.app.testing import setRoles, TEST_USER_ID
        from DateTime import DateTime
        portal = self.layer['portal']
        request = self.layer['request']
        setRoles(portal, TEST_USER_ID, ['Manager'])
        from bika.mining.idserver import REPORT_TYPE, DEFAULT_FORMAT
        setup = api.get_senaite_setup()
        rows = list(setup.getIDFormatting())
        for row in rows:
            if row['portal_type'] == REPORT_TYPE:
                row['form'] = DEFAULT_FORMAT
                row['sequence_type'] = 'generated'
                row['split_length'] = 1
        setup.setIDFormatting(rows)
        alsoProvides(request, IBikaMiningLayer)
        from plone.browserlayer.utils import registered_layers
        self.assertIn(IBikaMiningLayer, registered_layers())
        from archetypes.schemaextender.extender import disableCache
        disableCache(request)
        client = api.create(portal.clients, 'Client', title=u'Lotus test',
                            ClientID='LOTUS')
        sample = api.create(client, 'AnalysisRequest',
                            DateReceived=DateTime() - 1.0 / 24,
                            TargetTAT=8.0)
        sample.Schema().getField('TargetTAT').set(sample, 8.0)
        methods = api.get_senaite_setup().samplingdeviations
        method = next(obj for obj in methods.objectValues()
                      if obj.Title() == 'Composite')
        sample.Schema().getField('SamplingDeviation').set(sample, method)
        self.assertEqual(sample.Schema().getField('SamplingDeviation').widget.label,
                         'Sampling Method')
        workflow = api.get_tool('portal_workflow')
        chain = workflow.getChainFor(sample)
        workflow.setStatusOf(chain[0], sample, {'review_state': 'verified'})
        workflow[chain[0]].updateRoleMappingsFor(sample)
        sample.reindexObject()
        portal.uid_catalog.catalog_object(sample, sample._getURL())
        request.method = 'POST'
        view = MiningAjaxPublishView(sample, request)
        payload = {'template': REPORT_TEMPLATE, 'items': [api.get_uid(sample)],
                   'action': 'email', 'format': 'A4', 'orientation': 'portrait'}
        view.get_json = lambda: payload
        html = view.ajax_render_reports()
        self.assertIn('NO data', html)
        self.assertIn(api.get_uid(sample), html)
        payload['html'] = html
        preview = view.ajax_load_preview()
        self.assertIn('data:image/png', preview)
        url = view.ajax_save_reports()
        self.assertIn('/email?uids=', url)
        self.assertEqual(api.get_workflow_status_of(sample), 'verified')
        reports = [obj for obj in sample.objectValues()
                   if obj.portal_type == 'ResultsReport']
        self.assertEqual(len(reports), 1)
        self.assertEqual(reports[0].getPdf().contentType, 'application/pdf')
        self.assertTrue(reports[0].getPdf().data.startswith('%PDF'))
        self.assertEqual(reports[0].xlsx.contentType, MIME_TYPE)
        self.assertRegexpMatches(
            reports[0].xlsx.filename,
            r'^COA\d{2}-\d{4} LOTUS \d{6} \d{2}:\d{2}\.xlsx$')
        filename = reports[0].xlsx.filename
        saved_sheet = load_workbook(BytesIO(reports[0].xlsx.data))['Reporting']
        self.assertEqual(saved_sheet['A1'].value, 'Lotus test')
        self.assertEqual(saved_sheet['C1'].value, filename)
        self.assertIsNone(saved_sheet['E1'].value)
        self.assertIsNone(saved_sheet['G1'].value)
        self.assertEqual(api.normalize_filename(filename[:-5]),
                         api.get_id(reports[0]))
        self.assertEqual(reports[0].getContainedSamples(), [sample])
        workflow.doActionFor(sample, 'publish')
        refresh_xlsx(reports[0])
        self.assertEqual(reports[0].xlsx.filename, filename)
        data = reports[0].xlsx.data
        sheet = load_workbook(BytesIO(data))['Reporting']
        self.assertEqual(sheet['A4'].value, sample.getId())
        self.assertEqual(sheet['E4'].value, 'Composite')
        self.assertEqual(sheet['E1'].value,
                         sample.getDatePublished().strftime('%Y-%m-%d %H:%M'))
        self.assertTrue(sheet['G1'].value)
        self.assertEqual(sheet['M4'].value, 8.0)
        self.assertAlmostEqual(sheet['N4'].value, 1.0, places=2)
        self.assertEqual(sheet['O4'].value, 'On Time')
        self.assertIsNone(sheet['P4'].value)
        from bika.mining.browser.resultsxlsx import XLSXDownloadView
        downloaded = []
        request.response.write = downloaded.append
        XLSXDownloadView(reports[0], request)()
        self.assertEqual(downloaded, [data])
        self.assertEqual(request.response.getHeader('Content-Type'), MIME_TYPE)
        self.assertIn('.xlsx', request.response.getHeader('Content-Disposition'))

        from bika.mining.browser.resultsxlsx import ReportEmailView
        class Email(ReportEmailView):
            attachments = []
        Email.reports = reports
        email = Email(sample, request)
        self.assertEqual(email.generated_attachments, [reports[0].xlsx])
        attachments = email.email_attachments
        self.assertEqual(len(attachments), 1)
        self.assertEqual(attachments[0].get_filename(), filename)
        self.assertEqual(attachments[0].get_content_type(), MIME_TYPE)
        self.assertEqual(attachments[0].get_payload(decode=True), data)
        self.assertEqual(email.attachment_count, 1)
        self.assertAlmostEqual(email.total_size, len(data) / 1024.0)
        self.assertTrue(email.report_attachment_data[0]['url'].endswith('/download_xlsx'))

        from bika.mining.idserver import setup_xlsx_format, xlsx_filename
        setup = api.get_senaite_setup()
        rows = list(setup.getIDFormatting())
        for row in rows:
            if row['portal_type'] == REPORT_TYPE:
                row['form'] = u'TEST{year}-{seq:04d} {clientId} {hh:mm}'
        setup.setIDFormatting(rows)
        setup_xlsx_format()
        self.assertFalse(any(row['portal_type'] == 'MiningResultsXLSX'
                             for row in setup.getIDFormatting()))
        extra = api.create(sample, 'ResultsReport', sample=api.get_uid(sample))
        first = xlsx_filename(extra)
        self.assertEqual(xlsx_filename(extra), first)
        extra = api.create(sample, 'ResultsReport', sample=api.get_uid(sample))
        second = xlsx_filename(extra)
        self.assertRegexpMatches(first, r'^TEST\d{2}-0001 LOTUS \d{2}:\d{2}\.xlsx$')
        self.assertRegexpMatches(second, r'^TEST\d{2}-0002 LOTUS \d{2}:\d{2}\.xlsx$')

        from bika.mining.browser.resultsxlsx import DownloadReportsAdapter
        adapter = DownloadReportsAdapter(sample, request)
        adapter.download = lambda data, name, **kw: name
        self.assertEqual(adapter('download', [api.get_uid(reports[0])]), filename)


class TestMiningReportSelection(unittest.TestCase):
    def test_time_variable_uses_hours_and_minutes(self):
        from DateTime import DateTime
        from bika.mining.idserver import ReportTime
        self.assertEqual(u'{hh:mm}'.format(
            hh=ReportTime(DateTime('2026/10/08 09:15:00'))), u'09:15')

    def test_other_templates_keep_their_existing_handlers(self):
        from bika.mining.browser.resultsxlsx import MiningAjaxPublishView, REPORT_TEMPLATE
        calls = []
        class Legacy(object):
            def ajax_templates(self):
                return ['bika.coa:OtherMulti.pt']
            def ajax_render_reports(self, *args):
                calls.append('render')
                return 'normal report'
            def ajax_load_preview(self):
                calls.append('preview')
                return 'normal preview'
            def ajax_save_reports(self):
                calls.append('save')
                return 'normal save'
        view = MiningAjaxPublishView(None, None)
        view.legacy_view = lambda: Legacy()
        view.get_json = lambda: {'template': 'bika.coa:OtherMulti.pt'}
        self.assertEqual(view.ajax_templates(), ['bika.coa:OtherMulti.pt', REPORT_TEMPLATE])
        self.assertEqual(view.ajax_render_reports(), 'normal report')
        self.assertEqual(view.ajax_load_preview(), 'normal preview')
        self.assertEqual(view.ajax_save_reports(), 'normal save')
        self.assertEqual(calls, ['render', 'preview', 'save'])

    def test_mining_emails_exclude_placeholder_but_preserve_normal_pdf_and_csv(self):
        from bika.mining.browser.resultsxlsx import ReportEmailView, MIME_TYPE
        from ZPublisher.HTTPRequest import HTTPRequest
        from ZPublisher.HTTPResponse import HTTPResponse
        class Report(object):
            def __init__(self, pdf, xlsx=None, csv=None):
                self.pdf, self.xlsx, self.csv = pdf, xlsx, csv
            def getPdf(self):
                return self.pdf
            def getMetadata(self):
                return {'template': 'bika.coa:OtherMulti.pt'}
        def blob(name, kind, data):
            return type('Blob', (object,), {
                'data': data, 'filename': unicode(name), 'contentType': kind})()
        xlsx = blob('MiningResults.xlsx', MIME_TYPE, 'xlsx content')
        placeholder = blob('MiningResultsPreview.pdf', 'application/pdf', 'NO data')
        pdf = blob('OtherReport.pdf', 'application/pdf', 'normal PDF')
        csv = blob('OtherReport.csv', 'text/csv', 'normal CSV')
        class Email(ReportEmailView):
            reports = [Report(placeholder, xlsx), Report(placeholder, xlsx),
                       Report(pdf, csv=csv)]
            attachments = []
            email_csv_report_enabled = True
        request = HTTPRequest(BytesIO(), {'SERVER_NAME': 'localhost',
                                         'SERVER_PORT': '80',
                                         'REQUEST_METHOD': 'GET'}, HTTPResponse())
        view = Email(None, request)
        self.assertEqual(view.generated_attachments, [xlsx, pdf, csv])
        attachments = view.email_attachments
        self.assertEqual([a.get_filename() for a in attachments],
                         ['MiningResults.xlsx', 'OtherReport.pdf', 'OtherReport.csv'])
        self.assertEqual([a.get_content_type() for a in attachments],
                         [MIME_TYPE, 'application/pdf', 'text/csv'])

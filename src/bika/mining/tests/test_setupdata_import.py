# -*- coding: utf-8 -*-
import os
import unittest
from openpyxl import Workbook, load_workbook

from bika.mining.browser.overrides import setupdata


class Record(object):
    def __init__(self, **values):
        self.__dict__.update(values)

    def Title(self):
        return self.title

    def reindexObject(self):
        pass


class TestMatrixReferenceImport(unittest.TestCase):
    def setUp(self):
        self.created = []
        folder = Record(objectValues=lambda portal_type: self.created)
        self.importer = setupdata.Matrix_References(None)
        self.importer.context = Record(setup=Record(matrixreferences=folder))
        self.importer.lsd = Record(deferred=[])
        self.original_create = setupdata.api.create
        self.original_get_setup = setupdata.api.get_senaite_setup
        setupdata.api.get_senaite_setup = lambda: self.importer.context.setup

        def create(folder, portal_type, **values):
            self.assertIs(folder, self.importer.context.setup.matrixreferences)
            obj = Record(**values)
            self.created.append(obj)
            return obj

        setupdata.api.create = create

    def tearDown(self):
        setupdata.api.create = self.original_create
        setupdata.api.get_senaite_setup = self.original_get_setup

    def run_rows(self, rows):
        self.importer.get_rows = lambda start: iter(rows)
        self.importer.Import()

    def test_fields_and_deferred_references(self):
        self.run_rows([{"Matrix Ref": "Feed", "Area": "Plant",
                        "Circuit / Area": "Mill", "Target TAT (h)": 0,
                        "Plant Sample ID": "Feed point", "Sample Type": "Ore"}])
        obj = self.created[0]
        self.assertEqual(obj.title, "Feed")
        self.assertEqual(obj.area, "Plant")
        self.assertEqual(obj.circuit_area, "Mill")
        self.assertEqual(obj.target_tat, 0.0)
        deferred = self.importer.lsd.deferred
        self.assertEqual(len(deferred), 2)
        self.assertEqual(deferred[0]["src_field"], "plant_sample_id")
        self.assertEqual(deferred[0]["dest_query"]["title"], "Feed point")
        self.assertEqual(deferred[1]["dest_query"]["portal_type"], "SampleType")

    def test_duplicates_and_blank_titles(self):
        rows = [{"Title": "Feed"}, {"Title": " feed "}, {"Title": ""}]
        self.run_rows(rows)
        self.run_rows(rows)
        self.assertEqual(len(self.created), 1)
        self.assertIsNone(self.created[0].target_tat)

    def test_rejects_invalid_tat_before_creation(self):
        for value in ("invalid", "-1", "nan", "inf"):
            with self.assertRaises(ValueError):
                self.run_rows([{"Title": "Feed", "Target TAT": value}])
        self.assertEqual(self.created, [])

    def test_standard_worksheet_layout(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Matrix References"
        sheet.append(["Title", "Area", "Target TAT (h)"])
        sheet.append(["Description of columns", None, None])
        sheet.append(["Example only", None, None])
        sheet.append(["Feed", "Plant", 8])
        self.importer.lsd.context = self.importer.context
        self.importer(self.importer.lsd, workbook, "bika.mining", "test")
        self.assertEqual(len(self.created), 1)
        self.assertEqual(self.created[0].title, "Feed")
        self.assertEqual(self.created[0].target_tat, 8.0)

    def test_lotus_workbook(self):
        path = os.path.join(os.path.dirname(os.path.dirname(__file__)),
                            "setupdata", "Lotus new setup data.xlsx")
        workbook = load_workbook(path, data_only=True)
        sheet = workbook["Matrix References"]
        self.assertTrue(
            any(setupdata._key(cell.value) in
                ("title", "matrix ref", "matrix reference")
                for cell in next(sheet.iter_rows())),
            "Unrecognized Matrix References headers: %r" %
            ([cell.value for cell in next(sheet.iter_rows())],))
        expected = [tuple(cell.value for cell in row[:7])
                    for row in list(sheet.rows)[3:] if row[0].value]
        self.assertEqual(len(expected), 77)
        self.importer.lsd.context = self.importer.context
        self.importer(self.importer.lsd, workbook, "bika.mining", "Lotus")
        self.assertEqual(len(self.created), 77)
        deferred = self.importer.lsd.deferred
        self.assertEqual(len(deferred), 154)
        for index, values in enumerate(expected):
            title, area, circuit, point, description, tat, sample_type = values
            obj = self.created[index]
            self.assertEqual(obj.title, title)
            self.assertEqual(obj.area, area)
            self.assertEqual(obj.circuit_area, circuit)
            self.assertEqual(obj.target_tat, tat)
            self.assertIs(deferred[index * 2]["src_obj"], obj)
            self.assertEqual(deferred[index * 2]["dest_query"]["title"], point)
            self.assertEqual(deferred[index * 2 + 1]["dest_query"]["title"],
                             sample_type)
        self.assertEqual(self.created[0].title, "SM-01")
        self.assertEqual(self.created[1].title, "SM-02")
        self.assertEqual(self.created[1].target_tat, 0.0)
        self.importer(self.importer.lsd, workbook, "bika.mining", "Lotus")
        self.assertEqual(len(self.created), 77)
        self.assertEqual(len(deferred), 154)
